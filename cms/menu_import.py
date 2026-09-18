import csv
import io
import re
from datetime import date, timedelta

from django.utils import timezone

MEAL_LABELS = {
    "아침": "breakfast",
    "조식": "breakfast",
    "breakfast": "breakfast",
    "점심": "lunch",
    "중식": "lunch",
    "lunch": "lunch",
    "저녁": "dinner",
    "석식": "dinner",
    "dinner": "dinner",
    "간식": "snack",
    "snack": "snack",
}
WEEKDAYS = {
    "월": 0, "월요일": 0, "화": 1, "화요일": 1, "수": 2, "수요일": 2,
    "목": 3, "목요일": 3, "금": 4, "금요일": 4, "토": 5, "토요일": 5,
    "일": 6, "일요일": 6,
}


def _clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_date_cell(value, year=None):
    text = _clean(value)
    if not text:
        return None
    year = year or timezone.localdate().year
    embedded = re.search(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})|(\d{1,2})[./-](\d{1,2})", text)
    if embedded:
        groups = [g for g in embedded.groups() if g]
        try:
            if len(groups) == 3:
                return date(int(groups[0]), int(groups[1]), int(groups[2]))
            return date(year, int(groups[0]), int(groups[1]))
        except ValueError:
            pass
    normalized = text.replace("년", ".").replace("월", ".").replace("일", "")
    normalized = re.sub(r"\s+", "", normalized)
    match = re.match(r"^(\d{1,2})\.(\d{1,2})$", normalized)
    if match:
        try:
            return date(year, int(match.group(1)), int(match.group(2)))
        except ValueError:
            return None
    weekday = WEEKDAYS.get(re.sub(r"[^가-힣]", "", text))
    if weekday is not None:
        today = timezone.localdate()
        start = today - timedelta(days=today.weekday())
        return start + timedelta(days=weekday)
    return None


def _split_rows(text):
    rows = []
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if "\t" in line:
            cells = [c.strip() for c in line.split("\t")]
        elif "," in line and line.count(",") >= 2:
            cells = [c.strip() for c in next(csv.reader([line]))]
        else:
            cells = [line.strip()]
        if any(cells):
            rows.append(cells)
    return rows


def parse_menu_table(text, year=None):
    rows = _split_rows(text)
    if not rows:
        return []
    year = year or timezone.localdate().year

    if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
        header = rows[0]
        header_dates = [parse_date_cell(cell, year) for cell in header]
        if sum(1 for item in header_dates if item) >= 2:
            result = {}
            for row in rows[1:]:
                label = MEAL_LABELS.get(re.sub(r"\s+", "", row[0]) if row else "")
                if not label:
                    continue
                for idx, cell in enumerate(row[1:], start=1):
                    if idx >= len(header_dates) or not header_dates[idx] or not _clean(cell):
                        continue
                    day = result.setdefault(header_dates[idx], {})
                    day[label] = _clean(cell)
            return [{"date": day, **meals} for day, meals in sorted(result.items())]

        first_col_dates = [parse_date_cell(row[0], year) for row in rows[1:]]
        if sum(1 for item in first_col_dates if item) >= 1:
            labels = [MEAL_LABELS.get(re.sub(r"\s+", "", cell), "") for cell in header]
            result = []
            for row, day in zip(rows[1:], first_col_dates):
                if not day:
                    continue
                item = {"date": day}
                for idx, cell in enumerate(row):
                    if idx < len(labels) and labels[idx] and _clean(cell):
                        item[labels[idx]] = _clean(cell)
                if len(item) > 1:
                    result.append(item)
            if result:
                return result

    first = parse_date_cell(rows[0][0], year)
    dishes = []
    start = 1 if first else 0
    for row in rows[start:]:
        dishes.extend([cell for cell in row if _clean(cell)])
    if dishes:
        return [{"date": first or timezone.localdate(), "lunch": ", ".join(dishes)}]
    return []


def rows_from_xlsx(upload):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("엑셀 파일을 읽으려면 openpyxl이 필요합니다.") from exc
    workbook = load_workbook(upload, read_only=True, data_only=True)
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        cells = [_clean(cell) for cell in row]
        if any(cells):
            rows.append(cells)
    lines = ["\t".join(row) for row in rows]
    return parse_menu_table("\n".join(lines))


def rows_from_csv(upload):
    raw = upload.read()
    text = ""
    for encoding in ("utf-8-sig", "cp949", "utf-8"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        raise ValueError("CSV 글자 코드를 읽지 못했습니다.")
    return parse_menu_table(text)
